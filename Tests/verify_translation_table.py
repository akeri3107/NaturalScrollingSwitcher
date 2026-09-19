"""Read-only audit: python verify_translation_table.py workbook.xlsx catalog.xcstrings."""
import json
import re
import sys
import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
rows = list(workbook.active.values)[1:]
strings = json.load(open(sys.argv[2], encoding="utf-8"))["strings"]

def placeholders(value):
    names = list(dict.fromkeys(re.findall(r"\{([^}]+)\}", value)))
    return names

# User-approved menu wording supersedes this single original worksheet row.
menu_overrides = {"Natural Scrolling": ("NSS Control", "NSS 제어")}

for english, korean in rows:
    english, korean = menu_overrides.get(english, (english, korean))
    names = placeholders(english)
    def converted(value):
        for index, name in enumerate(names, 1):
            value = value.replace("{" + name + "}", f"%{index}$@")
        return value
    candidates = [entry for entry in strings.values()
                  if entry["localizations"]["en"]["stringUnit"]["value"] == converted(english)]
    assert candidates, english
    assert all(entry["localizations"]["ko"]["stringUnit"]["value"] == converted(korean)
               for entry in candidates), english
print(f"PASS: all {len(rows)} spreadsheet rows match (with approved menu override), including placeholders")
